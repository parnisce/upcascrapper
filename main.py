"""
Realtor.ca Public Listing Scraper & Drone Photography Lead Finder
built with Scrapling v0.5 for UPCA Real Estate Marketing & Media.
Extracts Agent Name, Phone Number, Email, Brokerage, and Drone Status.
"""

from __future__ import annotations

import argparse
import concurrent.futures
import csv
import hashlib
import logging
import mimetypes
import re
import sys
import time
import urllib.error
import urllib.parse
import urllib.request
from dataclasses import dataclass
from datetime import datetime, timezone
from pathlib import Path
from typing import Any, Dict, List, Optional, Set, Tuple

import openpyxl
import orjson
from openpyxl.styles import Alignment, Font, PatternFill
from openpyxl.utils import get_column_letter
from protego import Protego

from drone_analyzer import (
    LABEL_DRONE_DETECTED,
    LABEL_NO_DRONE,
    LABEL_UNCERTAIN,
    PhotoInspectionResult,
    analyze_single_photo,
    classify_listing_gallery,
    generate_html_review_report,
)
from history_store import ListingHistoryStore
from scrapling.engines.toolbelt.custom import Response
from scrapling.fetchers import FetcherSession
from scrapling.parser import Selector

BASE_URL = "https://www.realtor.ca"
ROBOTS_URL = f"{BASE_URL}/robots.txt"
TERMS_URL = f"{BASE_URL}/terms-of-use"
USER_AGENT = (
    "Mozilla/5.0 (Windows NT 10.0; Win64; x64) "
    "AppleWebKit/537.36 (KHTML, like Gecko) "
    "Chrome/131.0.0.0 Safari/537.36"
)

PROVINCE_SLUGS: Dict[str, str] = {
    "ontario": "on",
    "on": "on",
    "british columbia": "bc",
    "bc": "bc",
    "alberta": "ab",
    "ab": "ab",
    "quebec": "qc",
    "qc": "qc",
    "manitoba": "mb",
    "mb": "mb",
    "saskatchewan": "sk",
    "sk": "sk",
    "nova scotia": "ns",
    "ns": "ns",
    "new brunswick": "nb",
    "nb": "nb",
    "newfoundland and labrador": "nl",
    "nl": "nl",
    "prince edward island": "pe",
    "pe": "pe",
    "pei": "pe",
    "yukon": "yt",
    "yt": "yt",
    "northwest territories": "nt",
    "nt": "nt",
    "nunavut": "nu",
    "nu": "nu",
}

# Focused Lead Columns with Agent Contact Info
LEAD_COLUMNS: List[Tuple[str, str]] = [
    ("date_first_seen", "Date First Seen"),
    ("listing_date", "Published Listing Date"),
    ("property_address", "Property Address"),
    ("city", "City"),
    ("asking_price", "Asking Price"),
    ("agent_name", "Agent Name"),
    ("agent_phone", "Agent Phone Number"),
    ("agent_email", "Agent Email / Contact"),
    ("brokerage_name", "Brokerage"),
    ("brokerage_phone", "Brokerage Phone"),
    ("drone_photo_status", "Drone Photo Status"),
    ("marketing_lead_status", "Potential Marketing Lead Status"),
    ("num_accessible_photos", "Number of Photos"),
    ("detection_confidence", "Detection Confidence / Status"),
    ("classification_reason", "Reason for Classification"),
    ("listing_url", "Listing URL"),
    ("main_photo_url", "Main Photo URL"),
    ("last_checked_timestamp", "Last Checked Timestamp"),
]

FULL_COLUMNS: List[Tuple[str, str]] = [
    ("folder_id", "Folder ID"),
    ("listing_id", "Listing ID"),
    ("mls_number", "MLS Number"),
    ("is_new_label", "New vs Seen"),
    ("date_first_seen", "Date First Seen"),
    ("listing_date", "Published Listing Date"),
    ("property_address", "Property Address"),
    ("city", "City"),
    ("province", "Province"),
    ("asking_price", "Asking Price"),
    ("property_type", "Property Type"),
    ("bedrooms", "Bedrooms"),
    ("bathrooms", "Bathrooms"),
    ("interior_area", "Interior Area"),
    ("lot_size", "Lot Size"),
    ("agent_name", "Agent Name"),
    ("agent_phone", "Agent Phone Number"),
    ("agent_email", "Agent Email / Contact"),
    ("brokerage_name", "Brokerage"),
    ("brokerage_phone", "Brokerage Phone"),
    ("drone_photo_status", "Drone Photo Status"),
    ("marketing_lead_status", "Lead Status"),
    ("num_accessible_photos", "Number of Photos"),
    ("detection_confidence", "Confidence"),
    ("classification_reason", "Classification Reason"),
    ("listing_url", "Listing URL"),
    ("main_photo_url", "Main Photo URL"),
    ("local_image_folder", "Local Image Folder"),
    ("local_main_image_path", "Local Main Image Path"),
    ("collected_at", "Collected At"),
]


@dataclass
class ScraperConfig:
    city: str
    neighborhood: str
    postal_code_filter: str
    province: str
    province_slug: str
    city_slug: str
    search_url: str
    max_listings: int
    max_images_per_listing: int
    request_delay: float
    drone_filter: str  # "all", "no_drone", "has_drone", "uncertain"
    output_dir: Path
    images_dir: Path
    excel_path: Path
    csv_path: Path
    leads_csv_path: Path
    html_report_path: Path
    history_db_path: Path
    history_json_path: Path
    log_path: Path


def setup_logger(log_path: Path) -> logging.Logger:
    log_path.parent.mkdir(parents=True, exist_ok=True)
    logger = logging.getLogger("realtor_scraper")
    logger.setLevel(logging.INFO)

    # Preserve any in-memory dashboard handlers while replacing file/stdout handlers
    preserved = [
        h for h in logger.handlers if getattr(h, "is_dashboard_handler", False) or h.__class__.__name__ == "MemoryLogHandler"
    ]
    for h in list(logger.handlers):
        if h not in preserved:
            logger.removeHandler(h)
            try:
                h.close()
            except Exception:
                pass

    formatter = logging.Formatter(
        fmt="[%(asctime)s] %(levelname)s: %(message)s",
        datefmt="%Y-%m-%d %H:%M:%S",
    )

    file_handler = logging.FileHandler(log_path, mode="w", encoding="utf-8")
    file_handler.setFormatter(formatter)
    logger.addHandler(file_handler)

    console_handler = logging.StreamHandler(sys.stdout)
    console_handler.setFormatter(formatter)
    logger.addHandler(console_handler)

    return logger


def validate_config_and_location(
    location: str,
    max_listings: int = 10,
    max_images_per_listing: int = 50,
    request_delay: float = 1.0,
    drone_filter: str = "all",
    output_dir: Path = Path("output"),
) -> ScraperConfig:
    if not location or not location.strip():
        raise ValueError(
            "Location is required (e.g., 'Thorold, Ontario'). "
            "Pass --location 'City, Province' when running main.py."
        )

    parts = [p.strip() for p in location.split(",") if p.strip()]
    if len(parts) < 2:
        raise ValueError(
            f"Invalid location '{location}'. Please provide at least city and province "
            "(e.g., 'Thorold, Ontario' or 'Thorold Downtown, Thorold, Ontario')."
        )

    neighborhood = ""
    postal_code_filter = ""

    postal_match = re.search(r"\b([A-Za-z]\d[A-Za-z](?:\s*\d[A-Za-z]\d)?)\b", location)
    if postal_match:
        postal_code_filter = postal_match.group(1).upper().replace(" ", "")

    if len(parts) >= 3:
        neighborhood = parts[0]
        city_raw = parts[1]
        province_raw = parts[2].lower()
        province_display = parts[2].title()
    else:
        city_raw = parts[0]
        province_raw = parts[1].lower()
        province_display = parts[1].title()

    province_raw_clean = re.sub(r"\b[a-z]\d[a-z](?:\s*\d[a-z]\d)?\b", "", province_raw).strip()
    if province_raw_clean not in PROVINCE_SLUGS:
        raise ValueError(
            f"Unrecognized Canadian province '{province_display}'. "
            f"Supported provinces: {', '.join(sorted(set(PROVINCE_SLUGS.values())))}"
        )

    province_slug = PROVINCE_SLUGS[province_raw_clean]
    city_slug = re.sub(r"[^a-z0-9]+", "-", city_raw.lower()).strip("-")
    if not city_slug:
        raise ValueError(f"Could not derive a valid URL slug from city '{city_raw}'.")

    if max_listings < 1 or max_listings > 50:
        raise ValueError("max_listings must be between 1 and 50.")

    if max_images_per_listing < 1:
        raise ValueError("max_images_per_listing must be at least 1.")

    if neighborhood:
        neighborhood_slug = re.sub(r"[^a-z0-9]+", "-", neighborhood.lower()).strip("-")
        search_url = f"{BASE_URL}/{province_slug}/{city_slug}/{neighborhood_slug}/real-estate"
    else:
        search_url = f"{BASE_URL}/{province_slug}/{city_slug}/real-estate"

    images_dir = output_dir / "images"
    excel_path = output_dir / "realtor_listings.xlsx"
    csv_path = output_dir / "realtor_listings.csv"
    leads_csv_path = output_dir / "upca_drone_leads.csv"
    html_report_path = output_dir / "drone_review_report.html"
    history_db_path = output_dir / "listing_history.db"
    history_json_path = output_dir / "listing_history.json"
    log_path = output_dir / "scrape_log.txt"

    return ScraperConfig(
        city=city_raw.title(),
        neighborhood=neighborhood.title(),
        postal_code_filter=postal_code_filter,
        province=province_display,
        province_slug=province_slug,
        city_slug=city_slug,
        search_url=search_url,
        max_listings=max_listings,
        max_images_per_listing=max_images_per_listing,
        request_delay=max(0.5, request_delay),
        drone_filter=drone_filter,
        output_dir=output_dir,
        images_dir=images_dir,
        excel_path=excel_path,
        csv_path=csv_path,
        leads_csv_path=leads_csv_path,
        html_report_path=html_report_path,
        history_db_path=history_db_path,
        history_json_path=history_json_path,
        log_path=log_path,
    )


def load_robots_txt(session: FetcherSession, logger: logging.Logger) -> Optional[Protego]:
    try:
        logger.info("Checking robots.txt at %s", ROBOTS_URL)
        resp = session.get(ROBOTS_URL, stealthy_headers=False, timeout=15)
        if resp.status != 200:
            logger.warning("robots.txt returned HTTP %s", resp.status)
            return None
        robots_text = resp.text if isinstance(resp.text, str) else resp.body.decode("utf-8", errors="replace")
        return Protego.parse(robots_text)
    except Exception as exc:
        logger.error("Failed to fetch robots.txt: %s", exc)
        return None


def is_blocked_or_challenge_page(resp: Response) -> Tuple[bool, str]:
    if resp.status in (401, 403, 429, 503):
        return True, f"HTTP status {resp.status}"

    title_text = (resp.css("title::text").get() or "").strip().lower()
    if (
        "sorry, you have been blocked" in title_text
        or "access denied" in title_text
        or "just a moment" in title_text
        or "pardon our interruption" in title_text
    ):
        return True, f"Block/challenge page detected in title ('{title_text}')"

    return False, ""


def fetch_permitted_page(
    session: FetcherSession,
    url: str,
    robots_parser: Optional[Protego],
    logger: logging.Logger,
) -> Optional[Response]:
    if robots_parser is not None and not robots_parser.can_fetch(url, "*"):
        logger.warning("Disallowed by robots.txt (skipping): %s", url)
        return None

    try:
        resp = session.get(url, stealthy_headers=False, timeout=25)
    except Exception as exc:
        logger.error("HTTP request failed for %s: %s", url, exc)
        return None

    blocked, reason = is_blocked_or_challenge_page(resp)
    if blocked:
        logger.error(
            "Access restriction encountered on %s (%s). "
            "Stopping direct page access. Use official CREA DDF API for commercial feeds.",
            url,
            reason,
        )
        return None

    if resp.status != 200:
        logger.warning("Non-200 HTTP status (%s) for %s", resp.status, url)
        return None

    return resp


def extract_search_cards(search_resp: Response) -> List[Dict[str, str]]:
    results: List[Dict[str, str]] = []
    for card in search_resp.css("a.listingDetailsLink"):
        href = (card.attrib.get("href") or "").strip()
        if not href or "/real-estate/" not in href:
            continue

        full_url = search_resp.urljoin(href)
        mls_num = (card.css(".listingCardMLS span::text").get() or "").strip()
        price = (card.css(".listingCardPrice::text").get() or "").strip()
        address = (card.css(".listingCardAddress::text").get() or "").strip()
        office = (card.css(".listingCardOfficeName::text").get() or "").strip()
        card_img = (card.css("img.listingCardImage::attr(src)").get() or "").strip()

        bedrooms = ""
        bathrooms = ""
        sqft = ""
        for icon_con in card.css(".listingCardIconCon"):
            label = (icon_con.css(".listingCardIconText::text").get() or "").strip().lower()
            val = (icon_con.css(".listingCardIconNum::text").get() or "").strip()
            if "bedroom" in label:
                bedrooms = val
            elif "bathroom" in label:
                bathrooms = val
            elif "square" in label:
                sqft = val

        results.append(
            {
                "listing_url": full_url,
                "mls_number": mls_num,
                "asking_price": price,
                "full_address": address,
                "brokerage": office,
                "bedrooms": bedrooms,
                "bathrooms": bathrooms,
                "interior_area": sqft,
                "card_photo_url": card_img,
            }
        )
    return results


def _get_summary_value(resp: Selector, container_id: str) -> str:
    val = resp.css(f"#{container_id} .propertyDetailsSectionContentValue::text").get()
    if val:
        return str(val).strip()
    el = resp.css(f"#{container_id} .propertyDetailsSectionContentValue")
    if el:
        return el[0].get_all_text().strip()
    return ""


def extract_listing_data(
    detail_resp: Response,
    listing_url: str,
    fallback_card: Dict[str, str],
    logger: logging.Logger,
) -> Dict[str, Any]:
    product_ld: Dict[str, Any] = {}
    for script in detail_resp.css('script[type="application/ld+json"]'):
        raw_json = (script.text or "").strip()
        if not raw_json:
            continue
        try:
            parsed = orjson.loads(raw_json)
            if isinstance(parsed, dict) and parsed.get("@type") == "Product":
                product_ld = parsed
                break
        except Exception:
            continue

    url_match = re.search(r"/real-estate/(\d+)", listing_url)
    listing_id = url_match.group(1) if url_match else str(product_ld.get("sku") or "").strip()

    mls_els = detail_resp.css("#listingMLSNum")
    mls_text = str(mls_els[0].get_all_text()).strip() if mls_els else ""
    mls_match = re.search(r"Number:\s*([A-Za-z0-9-]+)", mls_text)
    mls_number = mls_match.group(1) if mls_match else fallback_card.get("mls_number", "")

    page_title = (detail_resp.css("title::text").get() or "").strip()
    listing_title = str(product_ld.get("name") or page_title or fallback_card.get("full_address", "")).strip()

    desc_els = detail_resp.css("#propertyDescriptionCon")
    dom_desc = str(desc_els[0].get_all_text()).strip() if desc_els else ""
    description_text = str(product_ld.get("description") or dom_desc or "").strip()

    street_address = (detail_resp.css("#listingAddressTitle::text").get() or "").strip()
    full_h1_html = detail_resp.css("#listingAddress").get() or ""
    city = ""
    province = ""

    if "<br>" in full_h1_html or "<br/>" in full_h1_html:
        parts = re.split(r"<br\s*/?>", full_h1_html, flags=re.IGNORECASE)
        first_line = Selector(parts[0]).get_all_text().strip()
        second_line = Selector(parts[1]).get_all_text().strip() if len(parts) > 1 else ""
        if not street_address:
            street_address = first_line
        if "," in second_line:
            city_part, prov_postal = second_line.split(",", 1)
            city = city_part.strip()
            prov_clean = re.sub(r"\b[A-Za-z]\d[A-Za-z]\s*\d[A-Za-z]\d\b", "", prov_postal).strip()
            province = prov_clean
    elif fallback_card.get("full_address"):
        addr_parts = [p.strip() for p in fallback_card["full_address"].split(",") if p.strip()]
        if len(addr_parts) >= 3:
            street_address = street_address or addr_parts[0]
            city = addr_parts[1]
            province = re.sub(r"\b[A-Za-z]\d[A-Za-z]\s*\d[A-Za-z]\d\b", "", addr_parts[2]).strip()

    asking_price = (
        (detail_resp.css("#listingPriceValue::text").get() or "").strip()
        or (detail_resp.css("#listingPrice::text").get() or "").strip()
        or fallback_card.get("asking_price", "")
    )

    prop_type = _get_summary_value(detail_resp, "propertyDetailsSectionContentSubCon_PropertyType")
    bldg_type = _get_summary_value(detail_resp, "propertyDetailsSectionContentSubCon_BuildingType")
    if prop_type and bldg_type and bldg_type.lower() not in prop_type.lower():
        property_type = f"{prop_type} - {bldg_type}"
    else:
        property_type = prop_type or bldg_type or str(product_ld.get("category") or "").replace(" > ", " - ").strip()

    bedrooms = (
        (detail_resp.css("#BedroomIcon .listingIconNum::text").get() or "").strip()
        or _get_summary_value(detail_resp, "propertyDetailsSectionVal_AboveGrade")
        or fallback_card.get("bedrooms", "")
    )
    bathrooms = (
        (detail_resp.css("#BathroomIcon .listingIconNum::text").get() or "").strip()
        or _get_summary_value(detail_resp, "propertyDetailsSectionVal_Total")
        or fallback_card.get("bathrooms", "")
    )

    interior_area = (
        _get_summary_value(detail_resp, "propertyDetailsSectionContentSubCon_SquareFootage")
        or (detail_resp.css("#SquareFootageIcon .listingIconNum::text").get() or "").strip()
        or fallback_card.get("interior_area", "")
    )
    if interior_area and "sqft" not in interior_area.lower() and "sq" not in interior_area.lower():
        interior_area = f"{interior_area} sqft"

    lot_size = _get_summary_value(detail_resp, "propertyDetailsSectionContentSubCon_LandSize")
    if not lot_size:
        frontage = _get_summary_value(detail_resp, "propertyDetailsSectionContentSubCon_Frontage")
        depth = _get_summary_value(detail_resp, "propertyDetailsSectionContentSubCon_LandDepth")
        if frontage and depth:
            lot_size = f"{frontage} x {depth}"
        elif frontage:
            lot_size = f"Frontage: {frontage}"

    # Agent Name, Agent Phone, Agent Email, Brokerage Name, Brokerage Phone
    agent_names = [
        (el.get_all_text() or "").strip()
        for el in detail_resp.css(".realtorCardName")
        if (el.get_all_text() or "").strip()
    ]
    agent_name = ", ".join(dict.fromkeys(agent_names))

    # Agent direct phone numbers
    agent_phones = [
        (el.get_all_text() or "").strip()
        for el in detail_resp.css(".realtorCardPhone .realtorCardContactNumber")
        if (el.get_all_text() or "").strip()
    ]
    agent_phone = ", ".join(dict.fromkeys(agent_phones))

    # Agent email check (Realtor.ca uses web contact forms, check for mailto or profile links)
    agent_email = ""
    for a in detail_resp.css(".realtorCardCon a[href^='mailto:'], .realtorCardBody a[href^='mailto:']"):
        href = a.attrib.get("href", "")
        if "mailto:" in href:
            agent_email = href.replace("mailto:", "").split("?")[0].strip()
            if agent_email:
                break
    if not agent_email:
        # Check text in description or page for email addresses
        email_matches = re.findall(r"\b[A-Za-z0-9._%+-]+@[A-Za-z0-9.-]+\.[A-Za-z]{2,}\b", detail_resp.body.decode("utf-8", errors="replace"))
        filtered_emails = [e for e in email_matches if "realtor.ca" not in e.lower() and "crea.ca" not in e.lower()]
        if filtered_emails:
            agent_email = filtered_emails[0]
        else:
            # Check for agent profile URL to route contact
            agent_profile_links = detail_resp.css(".realtorDetailsLink::attr(href)").getall()
            if agent_profile_links:
                agent_email = f"Realtor.ca Direct Form: {detail_resp.urljoin(agent_profile_links[0])}"
            else:
                agent_email = "Realtor.ca Web Form (Email Relay)"

    # Brokerage details
    office_names = [
        (el.get_all_text() or "").strip()
        for el in detail_resp.css(".officeCardName")
        if (el.get_all_text() or "").strip()
    ]
    brokerage_name = ", ".join(dict.fromkeys(office_names)) or fallback_card.get("brokerage", "")

    brokerage_phones = [
        (el.get_all_text() or "").strip()
        for el in detail_resp.css(".officeCardPhone .officeCardContactNumber")
        if (el.get_all_text() or "").strip()
    ]
    brokerage_phone = ", ".join(dict.fromkeys(brokerage_phones))

    if agent_name and brokerage_name:
        listing_agent_or_brokerage = f"{agent_name} ({brokerage_name})"
    else:
        listing_agent_or_brokerage = agent_name or brokerage_name

    listing_date = _get_summary_value(detail_resp, "propertyDetailsSectionContentSubCon_TimeOnRealtor")

    # Collect all accessible gallery photos
    photo_urls: List[str] = []
    seen_photos: Set[str] = set()

    def _add_photo(u: str) -> None:
        u_clean = u.strip()
        if not u_clean or "cdn.realtor.ca/listings/" not in u_clean:
            return
        basename = Path(urllib.parse.urlparse(u_clean).path).name.lower()
        dedup_key = basename or u_clean
        if dedup_key not in seen_photos:
            seen_photos.add(dedup_key)
            photo_urls.append(u_clean)

    for img in detail_resp.css("img"):
        src = img.attrib.get("src") or img.attrib.get("data-src") or ""
        if "highres" in src:
            _add_photo(detail_resp.urljoin(src))

    ld_images = product_ld.get("image")
    if isinstance(ld_images, list):
        for img_url in ld_images:
            if isinstance(img_url, str) and "highres" in img_url:
                _add_photo(img_url)
    for img in detail_resp.css("img"):
        src = img.attrib.get("src") or img.attrib.get("data-src") or ""
        _add_photo(detail_resp.urljoin(src))

    if not photo_urls and fallback_card.get("card_photo_url"):
        _add_photo(fallback_card["card_photo_url"])

    main_photo_url = photo_urls[0] if photo_urls else ""
    additional_photo_urls = photo_urls[1:] if len(photo_urls) > 1 else []

    now_iso = datetime.now(timezone.utc).astimezone().isoformat(timespec="seconds")

    logger.info(
        "Extracted listing %s (%s): %s | %s | Agent: %s (%s) | Photos: %d",
        listing_id or "unknown",
        mls_number or "no-MLS",
        street_address or listing_title,
        asking_price or "no-price",
        agent_name or "N/A",
        agent_phone or "N/A",
        len(photo_urls),
    )

    return {
        "listing_id": listing_id,
        "listing_url": listing_url,
        "mls_number": mls_number,
        "listing_title": listing_title,
        "description_text": description_text,
        "property_address": street_address,
        "city": city,
        "province": province,
        "asking_price": asking_price,
        "property_type": property_type,
        "bedrooms": bedrooms,
        "bathrooms": bathrooms,
        "interior_area": interior_area,
        "lot_size": lot_size,
        "agent_name": agent_name,
        "agent_phone": agent_phone,
        "agent_email": agent_email,
        "brokerage_name": brokerage_name,
        "brokerage_phone": brokerage_phone,
        "listing_agent_or_brokerage": listing_agent_or_brokerage,
        "listing_date": listing_date,
        "main_photo_url": main_photo_url,
        "additional_photo_urls_list": additional_photo_urls,
        "all_photo_urls": photo_urls,
        "num_accessible_photos": len(photo_urls),
        "collected_at": now_iso,
        "last_checked_timestamp": now_iso,
    }


def normalize_and_validate_fields(raw_data: Dict[str, Any], folder_id: str) -> Dict[str, Any]:
    def _clean(val: Any) -> str:
        if val is None:
            return ""
        text = re.sub(r"\s+", " ", str(val)).strip()
        if text.lower() in ("unknown", "n/a", "none", "null"):
            return ""
        return text

    main_photo = _clean(raw_data.get("main_photo_url", ""))
    add_photos_raw = raw_data.get("additional_photo_urls_list") or []
    add_photos: List[str] = []
    seen: Set[str] = {main_photo} if main_photo else set()
    for p in add_photos_raw:
        p_clean = _clean(p)
        if p_clean and p_clean not in seen:
            seen.add(p_clean)
            add_photos.append(p_clean)

    all_photos = ([main_photo] if main_photo else []) + add_photos

    return {
        "folder_id": folder_id,
        "listing_id": _clean(raw_data.get("listing_id")),
        "listing_url": _clean(raw_data.get("listing_url")),
        "mls_number": _clean(raw_data.get("mls_number")),
        "listing_title": _clean(raw_data.get("listing_title")),
        "description_text": _clean(raw_data.get("description_text")),
        "property_address": _clean(raw_data.get("property_address")),
        "city": _clean(raw_data.get("city")),
        "province": _clean(raw_data.get("province")),
        "asking_price": _clean(raw_data.get("asking_price")),
        "property_type": _clean(raw_data.get("property_type")),
        "bedrooms": _clean(raw_data.get("bedrooms")),
        "bathrooms": _clean(raw_data.get("bathrooms")),
        "interior_area": _clean(raw_data.get("interior_area")),
        "lot_size": _clean(raw_data.get("lot_size")),
        "agent_name": _clean(raw_data.get("agent_name")),
        "agent_phone": _clean(raw_data.get("agent_phone")),
        "agent_email": _clean(raw_data.get("agent_email")),
        "brokerage_name": _clean(raw_data.get("brokerage_name")),
        "brokerage_phone": _clean(raw_data.get("brokerage_phone")),
        "listing_agent_or_brokerage": _clean(raw_data.get("listing_agent_or_brokerage")),
        "listing_date": _clean(raw_data.get("listing_date")),
        "main_photo_url": main_photo,
        "additional_photo_urls_list": add_photos,
        "all_photo_urls": all_photos,
        "num_accessible_photos": len(all_photos),
        "additional_photo_urls": "; ".join(add_photos),
        "local_image_folder": "",
        "local_main_image_path": "",
        "additional_image_paths": "",
        "collected_at": _clean(raw_data.get("collected_at")),
        "last_checked_timestamp": _clean(raw_data.get("last_checked_timestamp")),
    }


def _determine_extension(content_type: str, url: str) -> str:
    ct_lower = content_type.split(";")[0].strip().lower()
    ext_map = {
        "image/jpeg": ".jpg",
        "image/jpg": ".jpg",
        "image/png": ".png",
        "image/webp": ".webp",
        "image/gif": ".gif",
        "image/avif": ".avif",
    }
    if ct_lower in ext_map:
        return ext_map[ct_lower]

    guessed = mimetypes.guess_extension(ct_lower)
    if guessed:
        return guessed

    parsed_path = urllib.parse.urlparse(url).path
    suffix = Path(parsed_path).suffix.lower()
    if suffix in (".jpg", ".jpeg", ".png", ".webp", ".gif"):
        return ".jpg" if suffix == ".jpeg" else suffix
    return ".jpg"


def _fetch_single_gallery_image(
    idx: int,
    img_url: str,
    listing_folder: Path,
    folder_id: str,
) -> Tuple[int, str, Optional[bytes], Optional[str], Optional[str]]:
    url_stem = Path(urllib.parse.urlparse(img_url).path).stem
    safe_stem = re.sub(r"[^A-Za-z0-9_-]+", "_", url_stem).strip("_")[:40] or f"img_{idx:02d}"

    existing_matches = list(listing_folder.glob(f"photo_{idx:02d}_{safe_stem}.*"))
    if existing_matches and existing_matches[0].stat().st_size > 0:
        file_path = existing_matches[0]
        img_bytes = file_path.read_bytes()
        rel_file_path = str(Path("output") / "images" / folder_id / file_path.name)
        return idx, img_url, img_bytes, rel_file_path, None

    req = urllib.request.Request(
        img_url,
        headers={
            "User-Agent": USER_AGENT,
            "Referer": BASE_URL + "/",
            "Accept": "image/avif,image/webp,image/apng,image/*,*/*;q=0.8",
        },
    )
    try:
        with urllib.request.urlopen(req, timeout=15) as resp:
            status = getattr(resp, "status", 200)
            content_type = resp.headers.get("Content-Type", "")
            if status != 200:
                return idx, img_url, None, None, f"HTTP {status}"
            if not content_type.lower().startswith("image/"):
                return idx, img_url, None, None, f"Invalid Content-Type '{content_type}'"
            img_bytes = resp.read()
            if not img_bytes:
                return idx, img_url, None, None, "Empty response body"

            ext = _determine_extension(content_type, img_url)
            filename = f"photo_{idx:02d}_{safe_stem}{ext}"
            file_path = listing_folder / filename
            file_path.write_bytes(img_bytes)
            rel_file_path = str(Path("output") / "images" / folder_id / filename)
            return idx, img_url, img_bytes, rel_file_path, None
    except Exception as exc:
        return idx, img_url, None, None, str(exc)


def download_and_inspect_gallery(
    listing: Dict[str, Any],
    images_base_dir: Path,
    max_images: int,
    logger: logging.Logger,
) -> Tuple[str, str, List[str], int, int, List[PhotoInspectionResult]]:
    folder_id = listing["folder_id"]
    listing_folder = images_base_dir / folder_id
    listing_folder.mkdir(parents=True, exist_ok=True)
    rel_folder = str(Path("output") / "images" / folder_id)

    candidate_urls = list(listing.get("all_photo_urls") or [])[:max_images]

    downloaded_paths: List[str] = []
    seen_hashes: Set[str] = set()
    photo_results: List[PhotoInspectionResult] = []
    success_count = 0
    failed_count = 0

    with concurrent.futures.ThreadPoolExecutor(max_workers=6) as pool:
        futures = [
            pool.submit(_fetch_single_gallery_image, idx, url, listing_folder, folder_id)
            for idx, url in enumerate(candidate_urls, start=1)
        ]
        ordered_outcomes = [f.result() for f in futures]

    ordered_outcomes.sort(key=lambda x: x[0])

    for idx, img_url, img_bytes, rel_file_path, err_msg in ordered_outcomes:
        if err_msg or not img_bytes or not rel_file_path:
            logger.warning("[%s] Image #%d failed (%s): %s", folder_id, idx, err_msg, img_url)
            failed_count += 1
            continue

        digest = hashlib.sha256(img_bytes).hexdigest()
        if digest in seen_hashes:
            continue
        seen_hashes.add(digest)

        downloaded_paths.append(rel_file_path)
        success_count += 1

        inspection = analyze_single_photo(
            img_bytes=img_bytes,
            index=idx,
            url=img_url,
            local_path=rel_file_path,
        )
        if inspection is not None:
            photo_results.append(inspection)

    main_local_path = downloaded_paths[0] if downloaded_paths else ""
    additional_local_paths = downloaded_paths[1:] if len(downloaded_paths) > 1 else []

    return (
        rel_folder,
        main_local_path,
        additional_local_paths,
        success_count,
        failed_count,
        photo_results,
    )


def _style_worksheet(
    ws: Any,
    columns: List[Tuple[str, str]],
    listings: List[Dict[str, Any]],
) -> None:
    header_font = Font(bold=True, color="FFFFFF")
    header_fill = PatternFill(start_color="0F172A", end_color="0F172A", fill_type="solid")

    fill_no_drone = PatternFill(start_color="DCFCE7", end_color="DCFCE7", fill_type="solid")
    fill_uncertain = PatternFill(start_color="FEF3C7", end_color="FEF3C7", fill_type="solid")
    fill_drone = PatternFill(start_color="FEE2E2", end_color="FEE2E2", fill_type="solid")

    headers = [header for _, header in columns]
    ws.append(headers)
    for col_idx in range(1, len(headers) + 1):
        cell = ws.cell(row=1, column=col_idx)
        cell.font = header_font
        cell.fill = header_fill
        cell.alignment = Alignment(vertical="center", wrap_text=True)

    status_col_idx: Optional[int] = None
    for idx, (key, _) in enumerate(columns, start=1):
        if key == "drone_photo_status":
            status_col_idx = idx
            break

    for row_idx, item in enumerate(listings, start=2):
        row_values = [str(item.get(key, "")) for key, _ in columns]
        ws.append(row_values)
        if status_col_idx is not None:
            status_cell = ws.cell(row=row_idx, column=status_col_idx)
            status_val = str(status_cell.value or "")
            if status_val == LABEL_NO_DRONE:
                status_cell.fill = fill_no_drone
            elif status_val == LABEL_UNCERTAIN:
                status_cell.fill = fill_uncertain
            elif status_val == LABEL_DRONE_DETECTED:
                status_cell.fill = fill_drone

    for col_idx, (_, header) in enumerate(columns, start=1):
        max_len = len(header)
        for row_idx in range(2, len(listings) + 2):
            cell_val = str(ws.cell(row=row_idx, column=col_idx).value or "")
            max_len = max(max_len, min(len(cell_val), 55))
        ws.column_dimensions[get_column_letter(col_idx)].width = max(max_len + 3, 14)

    ws.freeze_panes = "A2"


def export_to_excel_and_csv(
    listings: List[Dict[str, Any]],
    excel_path: Path,
    csv_path: Path,
    leads_csv_path: Path,
    logger: logging.Logger,
) -> None:
    excel_path.parent.mkdir(parents=True, exist_ok=True)

    # 1. Full CSV export
    with csv_path.open("w", newline="", encoding="utf-8-sig") as f:
        writer = csv.writer(f)
        writer.writerow([header for _, header in FULL_COLUMNS])
        for item in listings:
            writer.writerow([item.get(key, "") for key, _ in FULL_COLUMNS])
    logger.info("Saved full CSV export to %s (%d rows)", csv_path, len(listings))

    # 2. UPCA Drone Leads CSV (with Agent Name, Phone, Email)
    with leads_csv_path.open("w", newline="", encoding="utf-8-sig") as f:
        writer = csv.writer(f)
        writer.writerow([header for _, header in LEAD_COLUMNS])
        for item in listings:
            writer.writerow([item.get(key, "") for key, _ in LEAD_COLUMNS])
    logger.info("Saved UPCA Drone Leads CSV to %s (%d rows)", leads_csv_path, len(listings))

    # 3. Excel Workbook with 2 sheets
    wb = openpyxl.Workbook()
    ws_leads = wb.active
    ws_leads.title = "UPCA Drone Leads"
    _style_worksheet(ws_leads, LEAD_COLUMNS, listings)

    ws_full = wb.create_sheet(title="Full Listing Details")
    _style_worksheet(ws_full, FULL_COLUMNS, listings)

    leads_excel_path = excel_path.parent / "upca_drone_leads.xlsx"
    wb.save(leads_excel_path)
    logger.info("Saved UPCA Drone Leads Excel workbook to %s (%d rows)", leads_excel_path, len(listings))

    try:
        wb.save(excel_path)
        logger.info("Saved Excel spreadsheet to %s (%d rows)", excel_path, len(listings))
    except PermissionError:
        logger.warning(
            "Could not overwrite %s (currently open in Excel). "
            "Saved latest workbook to %s instead.",
            excel_path,
            leads_excel_path,
        )


def log_and_report_results(summary: Dict[str, Any], logger: logging.Logger) -> None:
    logger.info("=== UPCA Drone Lead & Scrape Run Summary ===")
    logger.info("Target Location               : %s, %s", summary["city"], summary["province"])
    logger.info("Search URL                    : %s", summary["search_url"])
    logger.info("Filter Applied                : %s", summary["drone_filter"])
    logger.info("Listings Collected            : %d", summary["listings_collected"])
    logger.info("  - Newly Discovered Listings : %d", summary["new_listings_count"])
    logger.info("  - Previously Seen Listings  : %d", summary["previously_seen_count"])
    logger.info("Drone Classification Breakdown:")
    logger.info("  - No drone photos detected  : %d", summary["no_drone_count"])
    logger.info("  - Uncertain (manual review) : %d", summary["uncertain_count"])
    logger.info("  - Drone photos detected     : %d", summary["drone_detected_count"])
    logger.info("Gallery Photos Inspected      : %d", summary["photos_downloaded"])
    logger.info("Excel Report                  : %s", summary["excel_path"])
    logger.info("UPCA Leads CSV                : %s", summary["leads_csv_path"])
    logger.info("HTML Visual Review Report     : %s", summary["html_report_path"])
    logger.info("History Database              : %s", summary["history_db_path"])


def run_scraper(
    location: str,
    max_listings: int = 10,
    max_images_per_listing: int = 50,
    request_delay: float = 1.0,
    drone_filter: str = "all",
    output_dir: Path = Path("output"),
) -> Dict[str, Any]:
    config = validate_config_and_location(
        location=location,
        max_listings=max_listings,
        max_images_per_listing=max_images_per_listing,
        request_delay=request_delay,
        drone_filter=drone_filter,
        output_dir=output_dir,
    )
    logger = setup_logger(config.log_path)
    history_store = ListingHistoryStore(config.history_db_path, config.history_json_path)

    logger.info(
        "Starting UPCA Realtor.ca scraper for '%s, %s' (filter=%s, max_listings=%d)",
        config.city,
        config.province,
        config.drone_filter,
        config.max_listings,
    )

    collected_listings: List[Dict[str, Any]] = []
    all_inspected_listings: List[Dict[str, Any]] = []
    seen_urls: Set[str] = set()
    seen_ids: Set[str] = set()
    seen_addresses: Set[str] = set()
    skipped_listings = 0
    total_photos_downloaded = 0
    total_photo_errors = 0

    with FetcherSession(impersonate="chrome", stealthy_headers=False, timeout=25) as session:
        robots_parser = load_robots_txt(session, logger)

        logger.info("Fetching public search page: %s ...", config.search_url)
        search_resp = fetch_permitted_page(session, config.search_url, robots_parser, logger)
        if search_resp is None:
            logger.error("Could not fetch search page %s. Aborting.", config.search_url)
            summary = {
                "city": config.city,
                "province": config.province,
                "search_url": config.search_url,
                "drone_filter": config.drone_filter,
                "listings_collected": 0,
                "listings_inspected": 0,
                "new_listings_count": 0,
                "previously_seen_count": 0,
                "no_drone_count": 0,
                "uncertain_count": 0,
                "drone_detected_count": 0,
                "listings_skipped": 1,
                "photos_downloaded": 0,
                "photo_errors": 0,
                "excel_path": str(config.excel_path),
                "csv_path": str(config.csv_path),
                "leads_csv_path": str(config.leads_csv_path),
                "html_report_path": str(config.html_report_path),
                "history_db_path": str(config.history_db_path),
                "log_path": str(config.log_path),
            }
            log_and_report_results(summary, logger)
            return summary

        cards = extract_search_cards(search_resp)
        logger.info("Found %d public listing links on %s", len(cards), config.search_url)

        for card_idx, card in enumerate(cards, start=1):
            if len(collected_listings) >= config.max_listings:
                break

            listing_url = card["listing_url"]
            url_id_match = re.search(r"/real-estate/(\d+)", listing_url)
            card_id = url_id_match.group(1) if url_id_match else ""
            card_street = (card.get("full_address") or "").split(",")[0].strip().lower()

            if (
                listing_url in seen_urls
                or (card_id and card_id in seen_ids)
                or (card_street and card_street in seen_addresses)
            ):
                logger.info(
                    "[%d/%d] Skipping duplicate listing (%s): %s",
                    card_idx,
                    len(cards),
                    card_street or card_id,
                    listing_url,
                )
                skipped_listings += 1
                continue

            logger.info(
                "[%d/%d] Fetching listing details: %s ...",
                card_idx,
                len(cards),
                card.get("full_address") or card_id or listing_url,
            )
            time.sleep(config.request_delay)
            detail_resp = fetch_permitted_page(session, listing_url, robots_parser, logger)
            if detail_resp is None:
                logger.warning("Skipping listing due to fetch error: %s", listing_url)
                skipped_listings += 1
                continue

            try:
                raw_data = extract_listing_data(detail_resp, listing_url, card, logger)
                lid = str(raw_data.get("listing_id") or "").strip()
                folder_id = f"listing_{lid}" if lid else f"listing_{len(all_inspected_listings) + 1:03d}"
                normalized = normalize_and_validate_fields(raw_data, folder_id)

                existing_record = history_store.lookup(
                    normalized["listing_id"],
                    normalized["listing_url"],
                )
                is_new_listing = existing_record is None

                logger.info(
                    "[%d/%d] Downloading & analyzing up to %d gallery photos for %s ...",
                    card_idx,
                    len(cards),
                    min(normalized["num_accessible_photos"], config.max_images_per_listing),
                    normalized["property_address"] or folder_id,
                )
                (
                    rel_folder,
                    main_img_path,
                    add_img_paths,
                    dl_ok,
                    dl_err,
                    photo_inspection_results,
                ) = download_and_inspect_gallery(
                    normalized,
                    config.images_dir,
                    config.max_images_per_listing,
                    logger,
                )
                normalized["local_image_folder"] = rel_folder
                normalized["local_main_image_path"] = main_img_path
                normalized["additional_image_paths"] = "; ".join(add_img_paths)

                gallery_analysis = classify_listing_gallery(
                    photo_results=photo_inspection_results,
                    total_accessible_photos=normalized["num_accessible_photos"],
                    description_text=normalized.get("description_text", ""),
                    is_new_listing=is_new_listing,
                )
                normalized["drone_photo_status"] = gallery_analysis.drone_photo_status
                normalized["detection_confidence"] = gallery_analysis.detection_confidence
                normalized["classification_reason"] = gallery_analysis.classification_reason
                normalized["marketing_lead_status"] = gallery_analysis.marketing_lead_status
                normalized["candidate_photo_indices"] = gallery_analysis.candidate_photo_indices
                normalized["photo_inspection_results"] = gallery_analysis.photo_results

                # Always persist every inspected listing into the history database so no data/contacts are lost
                date_first_seen, is_new_confirmed, times_seen = history_store.upsert_listing(
                    normalized,
                    normalized["last_checked_timestamp"],
                )
                normalized["date_first_seen"] = date_first_seen
                normalized["is_new_listing"] = is_new_confirmed
                normalized["is_new_label"] = (
                    "New Listing" if is_new_confirmed else f"Previously Seen ({times_seen}x)"
                )

                total_photos_downloaded += dl_ok
                total_photo_errors += dl_err

                seen_urls.add(listing_url)
                if normalized["listing_id"]:
                    seen_ids.add(normalized["listing_id"])
                if normalized["property_address"]:
                    seen_addresses.add(normalized["property_address"].lower())

                all_inspected_listings.append(normalized)

                # Filter condition check if user requested specific drone filter
                status = normalized["drone_photo_status"]
                if config.drone_filter == "no_drone" and status != LABEL_NO_DRONE:
                    logger.info(
                        "[%d/%d] Saved to DB but filtered from report (Status='%s', Filter='no_drone'): %s",
                        card_idx,
                        len(cards),
                        status,
                        normalized["property_address"],
                    )
                    continue
                elif config.drone_filter == "has_drone" and status != LABEL_DRONE_DETECTED:
                    logger.info(
                        "[%d/%d] Saved to DB but filtered from report (Status='%s', Filter='has_drone'): %s",
                        card_idx,
                        len(cards),
                        status,
                        normalized["property_address"],
                    )
                    continue
                elif config.drone_filter == "uncertain" and status != LABEL_UNCERTAIN:
                    logger.info(
                        "[%d/%d] Saved to DB but filtered from report (Status='%s', Filter='uncertain'): %s",
                        card_idx,
                        len(cards),
                        status,
                        normalized["property_address"],
                    )
                    continue

                logger.info(
                    "✅ [%s] %s | %s | Agent: %s (%s) | Status: %s | Lead: %s",
                    folder_id,
                    normalized["property_address"],
                    normalized["is_new_label"],
                    normalized["agent_name"],
                    normalized["agent_phone"],
                    normalized["drone_photo_status"],
                    normalized["marketing_lead_status"],
                )
                collected_listings.append(normalized)

            except Exception as exc:
                logger.error("Error processing listing %s: %s", listing_url, exc)
                skipped_listings += 1

    export_listings = collected_listings if collected_listings else all_inspected_listings
    if not collected_listings and all_inspected_listings:
        logger.warning(
            "0 of %d inspected listings matched strict filter '%s'. Exporting all %d inspected listings to reports so no data is lost.",
            len(all_inspected_listings),
            config.drone_filter,
            len(all_inspected_listings),
        )

    if export_listings:
        export_to_excel_and_csv(
            export_listings,
            config.excel_path,
            config.csv_path,
            config.leads_csv_path,
            logger,
        )
        generate_html_review_report(
            export_listings,
            config.html_report_path,
            f"{config.city}, {config.province}",
        )

    summary = {
        "city": config.city,
        "province": config.province,
        "search_url": config.search_url,
        "drone_filter": config.drone_filter,
        "listings_collected": len(collected_listings),
        "listings_inspected": len(all_inspected_listings),
        "new_listings_count": sum(1 for x in all_inspected_listings if x.get("is_new_listing")),
        "previously_seen_count": sum(1 for x in all_inspected_listings if not x.get("is_new_listing")),
        "no_drone_count": sum(
            1 for x in all_inspected_listings if x.get("drone_photo_status") == LABEL_NO_DRONE
        ),
        "uncertain_count": sum(
            1 for x in all_inspected_listings if x.get("drone_photo_status") == LABEL_UNCERTAIN
        ),
        "drone_detected_count": sum(
            1 for x in all_inspected_listings if x.get("drone_photo_status") == LABEL_DRONE_DETECTED
        ),
        "listings_skipped": skipped_listings,
        "photos_downloaded": total_photos_downloaded,
        "photo_errors": total_photo_errors,
        "excel_path": str(config.excel_path),
        "csv_path": str(config.csv_path),
        "leads_csv_path": str(config.leads_csv_path),
        "html_report_path": str(config.html_report_path),
        "history_db_path": str(config.history_db_path),
        "log_path": str(config.log_path),
    }
    log_and_report_results(summary, logger)
    return summary


def parse_args() -> argparse.Namespace:
    parser = argparse.ArgumentParser(
        description="Realtor.ca listing scraper & drone lead generator with agent contacts."
    )
    parser.add_argument(
        "--location",
        "-l",
        type=str,
        default="Thorold, Ontario",
        help="Target Canadian location (default: 'Thorold, Ontario')",
    )
    parser.add_argument(
        "--limit",
        "-n",
        type=int,
        default=10,
        help="Maximum listings to collect (default: 10)",
    )
    parser.add_argument(
        "--filter",
        "-f",
        choices=["all", "no_drone", "has_drone", "uncertain"],
        default="all",
        help="Filter listings by drone status ('all', 'no_drone', 'has_drone', 'uncertain')",
    )
    parser.add_argument(
        "--max-images",
        type=int,
        default=50,
        help="Maximum photos to inspect per listing (default: 50)",
    )
    parser.add_argument(
        "--delay",
        type=float,
        default=1.0,
        help="Request delay in seconds (default: 1.0)",
    )
    parser.add_argument(
        "--output-dir",
        type=Path,
        default=Path("output"),
        help="Output directory path (default: output)",
    )
    return parser.parse_args()


if __name__ == "__main__":
    args = parse_args()
    run_scraper(
        location=args.location,
        max_listings=args.limit,
        max_images_per_listing=args.max_images,
        request_delay=args.delay,
        drone_filter=args.filter,
        output_dir=args.output_dir,
    )
